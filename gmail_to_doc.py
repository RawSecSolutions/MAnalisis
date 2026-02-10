#!/usr/bin/env python3
"""
Gmail to Document Bot
Connects to Gmail via IMAP, fetches all emails from a specific sender,
and exports them to a Word (.docx) and/or PDF file.
"""

import imaplib
import email
import os
import sys
import argparse
from email.header import decode_header
from datetime import datetime
from dotenv import load_dotenv
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

load_dotenv()


def connect_to_gmail(email_address: str, app_password: str) -> imaplib.IMAP4_SSL:
    """Connect to Gmail via IMAP and return the connection."""
    imap = imaplib.IMAP4_SSL("imap.gmail.com")
    imap.login(email_address, app_password)
    return imap


def decode_mime_header(header_value: str) -> str:
    """Decode a MIME-encoded email header into a readable string."""
    if header_value is None:
        return ""
    decoded_parts = decode_header(header_value)
    result = []
    for part, charset in decoded_parts:
        if isinstance(part, bytes):
            result.append(part.decode(charset or "utf-8", errors="replace"))
        else:
            result.append(part)
    return " ".join(result)


def extract_body(msg: email.message.Message) -> str:
    """Extract the plain text body from an email message."""
    body = ""
    if msg.is_multipart():
        for part in msg.walk():
            content_type = part.get_content_type()
            content_disposition = str(part.get("Content-Disposition", ""))
            if content_type == "text/plain" and "attachment" not in content_disposition:
                charset = part.get_content_charset() or "utf-8"
                payload = part.get_payload(decode=True)
                if payload:
                    body += payload.decode(charset, errors="replace")
    else:
        content_type = msg.get_content_type()
        if content_type == "text/plain":
            charset = msg.get_content_charset() or "utf-8"
            payload = msg.get_payload(decode=True)
            if payload:
                body = payload.decode(charset, errors="replace")
    return body.strip()


def fetch_emails(imap: imaplib.IMAP4_SSL, sender_email: str, folder: str = "INBOX") -> list[dict]:
    """Fetch all emails from a specific sender and return them as a list of dicts."""
    imap.select(folder, readonly=True)

    status, message_ids = imap.search(None, f'(FROM "{sender_email}")')
    if status != "OK" or not message_ids[0]:
        print(f"No emails found from {sender_email}")
        return []

    ids = message_ids[0].split()
    print(f"Found {len(ids)} email(s) from {sender_email}")

    emails = []
    for i, msg_id in enumerate(ids, 1):
        status, msg_data = imap.fetch(msg_id, "(RFC822)")
        if status != "OK":
            continue

        raw_email = msg_data[0][1]
        msg = email.message_from_bytes(raw_email)

        subject = decode_mime_header(msg["Subject"])
        from_addr = decode_mime_header(msg["From"])
        date_str = decode_mime_header(msg["Date"])
        body = extract_body(msg)

        try:
            date_parsed = email.utils.parsedate_to_datetime(date_str)
            date_formatted = date_parsed.strftime("%Y-%m-%d %H:%M:%S")
        except Exception:
            date_formatted = date_str

        emails.append({
            "number": i,
            "subject": subject,
            "from": from_addr,
            "date": date_formatted,
            "body": body,
        })
        print(f"  [{i}/{len(ids)}] Fetched: {subject[:60]}")

    # Sort by date ascending
    emails.sort(key=lambda e: e["date"])
    return emails


def export_to_docx(emails: list[dict], sender_email: str, output_path: str):
    """Export emails to a Word (.docx) document."""
    doc = Document()

    # Title
    title = doc.add_heading(f"Emails from: {sender_email}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total: {len(emails)} emails")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(128, 128, 128)

    doc.add_paragraph()  # Spacer

    for i, em in enumerate(emails):
        # Email header
        heading = doc.add_heading(f"Email #{em['number']}: {em['subject'] or '(No Subject)'}", level=2)

        # Date and from info
        info_para = doc.add_paragraph()
        date_run = info_para.add_run(f"Date: {em['date']}")
        date_run.font.size = Pt(9)
        date_run.font.color.rgb = RGBColor(100, 100, 100)
        info_para.add_run("\n")
        from_run = info_para.add_run(f"From: {em['from']}")
        from_run.font.size = Pt(9)
        from_run.font.color.rgb = RGBColor(100, 100, 100)

        # Body
        if em["body"]:
            body_para = doc.add_paragraph()
            body_run = body_para.add_run(em["body"])
            body_run.font.size = Pt(11)
        else:
            no_body = doc.add_paragraph()
            run = no_body.add_run("(No plain text content)")
            run.font.italic = True
            run.font.color.rgb = RGBColor(150, 150, 150)

        # Separator between emails
        if i < len(emails) - 1:
            separator = doc.add_paragraph()
            separator.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sep_run = separator.add_run("─" * 50)
            sep_run.font.color.rgb = RGBColor(200, 200, 200)

    doc.save(output_path)
    print(f"Word document saved: {output_path}")


def export_to_pdf(emails: list[dict], sender_email: str, output_path: str):
    """Export emails to a PDF document."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Title
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 12, f"Emails from: {sender_email}", new_x="LMARGIN", new_y="NEXT", align="C")

    # Metadata
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(128, 128, 128)
    pdf.cell(0, 8, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total: {len(emails)} emails",
             new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(5)

    for i, em in enumerate(emails):
        pdf.set_text_color(0, 0, 0)

        # Subject
        pdf.set_font("Helvetica", "B", 13)
        subject_text = f"Email #{em['number']}: {em['subject'] or '(No Subject)'}"
        pdf.multi_cell(0, 7, subject_text, new_x="LMARGIN", new_y="NEXT")

        # Date & from
        pdf.set_font("Helvetica", "", 8)
        pdf.set_text_color(100, 100, 100)
        pdf.cell(0, 5, f"Date: {em['date']}", new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 5, f"From: {em['from']}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)

        # Body
        pdf.set_text_color(0, 0, 0)
        if em["body"]:
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(0, 5, em["body"], new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Helvetica", "I", 10)
            pdf.set_text_color(150, 150, 150)
            pdf.cell(0, 5, "(No plain text content)", new_x="LMARGIN", new_y="NEXT")

        # Separator
        if i < len(emails) - 1:
            pdf.ln(5)
            pdf.set_draw_color(200, 200, 200)
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.w - pdf.r_margin, pdf.get_y())
            pdf.ln(5)

    pdf.output(output_path)
    print(f"PDF saved: {output_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Fetch all Gmail emails from a specific sender and export to Word/PDF."
    )
    parser.add_argument(
        "sender",
        help="Email address of the sender to search for (e.g. john@example.com)",
    )
    parser.add_argument(
        "--format",
        choices=["docx", "pdf", "both"],
        default="both",
        help="Output format: docx, pdf, or both (default: both)",
    )
    parser.add_argument(
        "--output",
        default="emails",
        help="Output filename without extension (default: 'emails')",
    )
    parser.add_argument(
        "--folder",
        default="INBOX",
        help="Gmail folder to search (default: INBOX). Use '[Gmail]/All Mail' for all mail.",
    )
    args = parser.parse_args()

    # Load credentials from environment
    gmail_user = os.getenv("GMAIL_ADDRESS")
    gmail_password = os.getenv("GMAIL_APP_PASSWORD")

    if not gmail_user or not gmail_password:
        print("Error: Set GMAIL_ADDRESS and GMAIL_APP_PASSWORD in your .env file.")
        print("See .env.example for details.")
        sys.exit(1)

    # Connect
    print(f"Connecting to Gmail as {gmail_user}...")
    try:
        imap = connect_to_gmail(gmail_user, gmail_password)
    except imaplib.IMAP4.error as e:
        print(f"Login failed: {e}")
        print("Make sure you're using an App Password, not your regular password.")
        sys.exit(1)

    # Fetch emails
    print(f"Searching for emails from: {args.sender}")
    emails = fetch_emails(imap, args.sender, folder=args.folder)
    imap.logout()

    if not emails:
        print("No emails found. Nothing to export.")
        sys.exit(0)

    # Export
    if args.format in ("docx", "both"):
        docx_path = f"{args.output}.docx"
        export_to_docx(emails, args.sender, docx_path)

    if args.format in ("pdf", "both"):
        pdf_path = f"{args.output}.pdf"
        export_to_pdf(emails, args.sender, pdf_path)

    print("Done!")


if __name__ == "__main__":
    main()
